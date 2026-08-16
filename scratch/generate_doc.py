import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=140, bottom=140, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_word_manual(output_path):
    doc = docx.Document()

    # Set page margins (1 inch = 72 pt)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("Talent Intelligence Engine (TIE)")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(16, 42, 67) # Deep Navy

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_sub.add_run("Enterprise Operational Workflow & Complete User Manual")
    run_sub.font.name = 'Calibri'
    run_sub.font.size = Pt(16)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(42, 107, 107) # Teal

    doc.add_paragraph() # Spacer

    # Callout Box: Executive Intro
    table_callout = doc.add_table(rows=1, cols=1)
    table_callout.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table_callout.rows[0].cells[0]
    cell.width = Inches(6.5)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=160, bottom=160, left=200, right=200)

    p_call = cell.paragraphs[0]
    r_call = p_call.add_run("Welcome to the Talent Intelligence Engine (TIE) Enterprise Platform Manual. This operational guide outlines the complete step-by-step workflow across all 4 system roles: Super Admin, HR Admin, Team Manager, and Employee/Team Member.")
    r_call.font.name = 'Calibri'
    r_call.font.size = Pt(11)
    r_call.font.italic = True
    r_call.font.color.rgb = RGBColor(36, 59, 83)

    doc.add_paragraph()

    # Section 1: Executive Architecture & Role Matrix
    h1 = doc.add_heading("1. Executive Architecture & Role Matrix", level=1)
    h1.runs[0].font.color.rgb = RGBColor(16, 42, 67)

    p = doc.add_paragraph()
    r = p.add_run("The TIE platform operates on a 4-tier hierarchical access matrix, enforcing database-level tenant insulation (Row Level Security) and strict permission boundaries:")
    r.font.name = 'Calibri'

    # Table for Role Matrix
    table = doc.add_table(rows=5, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Role", "Primary Responsibilities", "Access Level"]
    col_widths = [Inches(1.5), Inches(3.2), Inches(1.8)]

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        set_cell_background(hdr_cells[i], "102A43")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        run = p.add_run(title)
        run.font.name = 'Calibri'
        run.font.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ("1. Super Admin", "System oversight, tenant provisioning, company creation, HR Admin activation", "Platform-Wide Control"),
        ("2. HR Admin", "Enterprise workspace management, team creation, Manager invitation, org-wide reporting", "Company-Wide Control"),
        ("3. Team Manager", "Team member onboarding (CSV / Email), assessment tracking, report review, 30-Day Plan generation", "Team-Wide Control"),
        ("4. Employee", "Secure onboarding, completion of behavioral assessment, individual report access", "Personal Portal Control")
    ]

    for row_idx, row_data in enumerate(data, start=1):
        row_cells = table.rows[row_idx].cells
        fill_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            row_cells[col_idx].width = col_widths[col_idx]
            set_cell_background(row_cells[col_idx], fill_color)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[col_idx].paragraphs[0]
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            if col_idx == 0:
                run.font.bold = True
                run.font.color.rgb = RGBColor(42, 107, 107)

    doc.add_paragraph()

    # Section 2: Single-Use Invitation & Security Lifecycle
    h2 = doc.add_heading("2. Single-Use Invitation & Security Lifecycle", level=1)
    h2.runs[0].font.color.rgb = RGBColor(16, 42, 67)

    rules = [
        ("Single-Use Link Validity: ", "Every activation link emailed to HR Admins, Managers, and Employees is cryptographically locked to that specific user email and role. Once clicked and used to set a password, the token status changes to 'accepted' and expires instantly after 1 use."),
        ("Resend & Recovery Protocol: ", "If an invite link expires, is misplaced, or fails, the Admin or Manager simply clicks 'Invite' / 'Resend' again for that email. The backend automatically revokes/deletes the old pending token and dispatches a fresh activation URL."),
        ("Automatic Session Cleansing: ", "Opening an invitation link automatically clears any previous browser sessions or cached cookies to prevent account overlap.")
    ]

    for title, desc in rules:
        p = doc.add_paragraph(style='List Bullet')
        r1 = p.add_run(title)
        r1.font.name = 'Calibri'
        r1.font.bold = True
        r1.font.color.rgb = RGBColor(42, 107, 107)
        r2 = p.add_run(desc)
        r2.font.name = 'Calibri'

    doc.add_paragraph()

    # Section 3: Detailed Step-by-Step Role Workflows
    doc.add_heading("3. Detailed Step-by-Step Role Workflows", level=1).runs[0].font.color.rgb = RGBColor(16, 42, 67)

    # Phase 1
    doc.add_heading("Phase 1: Super Admin Workflow (Tenant Provisioning)", level=2).runs[0].font.color.rgb = RGBColor(42, 107, 107)

    steps_p1 = [
        ("Step 1.1: Log In to Super Admin Control Panel", "Navigate to /login and enter Super Admin credentials. System redirects to /super-admin."),
        ("Step 1.2: Register a New Client Company", "Under 'Company Management', click 'Provision New Client Company'. Enter Company Name (e.g. Acme Corp) and Description. System auto-generates corporate slug (/acme-corp)."),
        ("Step 1.3: Invite HR Admin", "Under 'HR Admin Invitation Panel', select the target company. Enter HR Admin Email, First Name, and Last Name. Click 'Send HR Activation Link'. If email bounces, click 'Resend'.")
    ]

    for title, text in steps_p1:
        p = doc.add_paragraph()
        r1 = p.add_run(title + "\n")
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(16, 42, 67)
        r2 = p.add_run(text)
        r2.font.size = Pt(10.5)

    # Phase 2
    doc.add_heading("Phase 2: HR Admin Workflow (Enterprise Setup & Team Building)", level=2).runs[0].font.color.rgb = RGBColor(42, 107, 107)

    steps_p2 = [
        ("Step 2.1: Activate HR Admin Account", "Open email inbox and click 'Activate Your HR Admin Account'. On /accept-invite?token=..., fill in First/Last Name, Designation, Experience, and set Password. Click 'Complete Activation'."),
        ("Step 2.2: Access Corporate HR Dashboard", "Log in to access /[company_slug]/hr_admin to monitor company-wide metrics and teams."),
        ("Step 2.3: Create Organizational Teams", "Click 'Create New Team'. Enter Team Name (e.g. Tech Team, Marketing Team) and click 'Provision Team'."),
        ("Step 2.4: Invite Team Managers / Leads", "Click 'Invite Manager'. Select destination team, enter Manager's Email, First/Last Name, and click 'Dispatch Manager Invitation'.")
    ]

    for title, text in steps_p2:
        p = doc.add_paragraph()
        r1 = p.add_run(title + "\n")
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(16, 42, 67)
        r2 = p.add_run(text)
        r2.font.size = Pt(10.5)

    # Phase 3
    doc.add_heading("Phase 3: Team Manager Workflow (Onboarding & Plan Generation)", level=2).runs[0].font.color.rgb = RGBColor(42, 107, 107)

    steps_p3 = [
        ("Step 3.1: Activate Manager Account", "Manager opens activation email, clicks single-use link, fills account details, and accesses /[company_slug]/[team_slug]/manager."),
        ("Step 3.2: Onboard Team Members (2 Flexible Methods)", "Managers can add members using Individual Email Onboarding or CSV Bulk Upload.\nCSV Format Required:\nemail,first_name,last_name,designation,experience_years\njohn@company.com,John,Doe,Software Engineer,5"),
        ("Step 3.3: Review Behavioral Reports", "As employees complete assessment, status updates live. Click 'View Behavioral Report' to review Work Preference Analysis (Workplace Impact, Early Risk Indicators, Thrive Conditions)."),
        ("Step 3.4: Generate 30-Day Executive Action Plan", "Click 'Generate 30-Day Plan' inside the report preview. The AI synthesizes a 4-week roadmap with employee commitments, manager coaching support, and review checkpoints. Click 'Download Plan PDF' for presentation.")
    ]

    for title, text in steps_p3:
        p = doc.add_paragraph()
        r1 = p.add_run(title + "\n")
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(16, 42, 67)
        r2 = p.add_run(text)
        r2.font.size = Pt(10.5)

    # Phase 4
    doc.add_heading("Phase 4: Employee / Team Member Workflow (Assessment & Output)", level=2).runs[0].font.color.rgb = RGBColor(42, 107, 107)

    steps_p4 = [
        ("Step 4.1: Activate Employee Account", "Employee opens email, clicks 'Take Assessment & Activate Account', and creates password on /accept-invite."),
        ("Step 4.2: Complete Behavioral Assessment", "Navigates to /[company_slug]/[team_slug]/user, clicks 'Start Assessment', and completes scenario questions."),
        ("Step 4.3: View Personal Executive Report & Plan", "Instantly gains access to Work Preference Analysis Report and assigned 30-Day Action Plan.")
    ]

    for title, text in steps_p4:
        p = doc.add_paragraph()
        r1 = p.add_run(title + "\n")
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(16, 42, 67)
        r2 = p.add_run(text)
        r2.font.size = Pt(10.5)

    # Section 4: Quick Reference Table
    doc.add_heading("4. Quick Reference: Operational Scenarios", level=1).runs[0].font.color.rgb = RGBColor(16, 42, 67)

    ref_table = doc.add_table(rows=5, cols=2)
    ref_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    ref_table.autofit = False

    ref_headers = ["Issue / Scenario", "Recommended Action"]
    ref_widths = [Inches(2.5), Inches(4.0)]

    for i, title in enumerate(ref_headers):
        cell = ref_table.rows[0].cells[i]
        cell.width = ref_widths[i]
        set_cell_background(cell, "102A43")
        set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.font.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)

    ref_data = [
        ("User missed or lost activation email", "Go to Invitations panel, click 'Resend Invite' next to email. Old link is revoked and new link sent."),
        ("Invite link says 'Invalid or Expired'", "The link was already used to set up account. Direct user to /login to sign in with password."),
        ("Manager needs to onboard 50+ members", "Use CSV Bulk Upload in Manager Dashboard to invite all members in 1 click."),
        ("Need PDF for client handover", "Click 'Export as PDF' on any Report or Action Plan. Document renders with clean 15mm top/bottom page margins.")
    ]

    for row_idx, row_data in enumerate(ref_data, start=1):
        cells = ref_table.rows[row_idx].cells
        fill_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cells[col_idx].width = ref_widths[col_idx]
            set_cell_background(cells[col_idx], fill_color)
            set_cell_margins(cells[col_idx], top=100, bottom=100, left=140, right=140)
            p = cells[col_idx].paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(10)

    # Footer
    doc.add_paragraph()
    p_foot = doc.add_paragraph()
    p_foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_foot = p_foot.add_run("Talent Intelligence Engine (TIE) • Document Version 2.0 (Enterprise Edition)")
    r_foot.font.size = Pt(9)
    r_foot.font.italic = True
    r_foot.font.color.rgb = RGBColor(148, 163, 184)

    doc.save(output_path)
    print(f"Successfully created Word document at: {output_path}")

if __name__ == "__main__":
    out_dir = r"C:\Users\kjsan\OneDrive\Documents\TIE Project\TIE"
    out_path = os.path.join(out_dir, "Talent_Intelligence_Engine_Enterprise_User_Manual.docx")
    create_word_manual(out_path)
